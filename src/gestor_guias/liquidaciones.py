"""Liquidaciones: pago semanal de contratistas y liquidacion laboral definitiva.

Dos cosas distintas, en modulos separados del panel:

- **Semanal (contrato de SERVICIOS)**: al contratista se le paga un valor
  fijo por cada encomienda entregada (estado `E`) durante la semana, menos
  las cuotas de prestamos o adelantos que tenga pendientes.
- **Laboral definitiva (contrato de NOMINA)**: al retirar a un empleado se
  liquidan cesantias, intereses de cesantias, prima y vacaciones segun su
  fecha de ingreso, mas la indemnizacion si aplica.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .nomina import CURRENCY_FORMAT, HEADER_FILL, HEADER_FONT, TITLE_FONT, descuento_nomina_empleado
from .repository import GuiaRepository


CONTRATO_NOMINA = "NOMINA"
CONTRATO_SERVICIOS = "SERVICIOS"
CONTRATOS_VALIDOS = (CONTRATO_NOMINA, CONTRATO_SERVICIOS)

# Convenciones laborales colombianas para la liquidacion definitiva.
DIAS_ANIO_LABORAL = 360
TASA_INTERESES_CESANTIAS = 0.12
DIAS_VACACIONES_BASE = 720  # 15 dias habiles por anio trabajado


def inicio_de_semana(fecha: str | date) -> date:
    """Lunes de la semana a la que pertenece la fecha."""
    dia = date.fromisoformat(str(fecha)[:10]) if isinstance(fecha, str) else fecha
    return dia - timedelta(days=dia.weekday())


def rango_semana(fecha: str | date) -> tuple[date, date]:
    """(lunes, domingo) de la semana de la fecha."""
    lunes = inicio_de_semana(fecha)
    return lunes, lunes + timedelta(days=6)


def dias_laborales_360(desde: str, hasta: str) -> int:
    """Dias entre dos fechas con la convencion de 360 dias (meses de 30).

    Es la base que usa la liquidacion laboral en Colombia.
    """
    inicio = date.fromisoformat(str(desde)[:10])
    fin = date.fromisoformat(str(hasta)[:10])
    if fin < inicio:
        return 0
    dias_inicio = min(inicio.day, 30)
    dias_fin = min(fin.day, 30)
    return (fin.year - inicio.year) * 360 + (fin.month - inicio.month) * 30 + (dias_fin - dias_inicio)


def calcular_liquidacion_semanal(
    entregas: int,
    valor_encomienda: int,
    descuento_prestamos: int = 0,
    otros_descuentos: int = 0,
) -> dict:
    """Pago semanal de un contratista por encomiendas entregadas."""
    entregas = max(0, int(entregas or 0))
    valor_encomienda = int(valor_encomienda or 0)
    subtotal = entregas * valor_encomienda
    descuentos = int(descuento_prestamos or 0) + int(otros_descuentos or 0)
    return {
        "entregas": entregas,
        "valor_encomienda": valor_encomienda,
        "subtotal": subtotal,
        "descuento_prestamos": int(descuento_prestamos or 0),
        "otros_descuentos": int(otros_descuentos or 0),
        "total_pagar": subtotal - descuentos,
    }


def preparar_semana(repository: GuiaRepository, fecha: str) -> dict:
    """Datos de la semana para liquidar a todos los contratistas de servicios."""
    lunes, domingo = rango_semana(fecha)
    semana_inicio = lunes.isoformat()

    entregas_por_operador = repository.contar_entregas_periodo(semana_inicio, domingo.isoformat())
    guardadas = {
        registro["empleado"]: registro
        for registro in repository.listar_liquidaciones_semanales(semana_inicio)
    }
    periodo_mes = semana_inicio[:7]

    empleados = []
    for empleado in repository.listar_empleados_nomina(CONTRATO_SERVICIOS):
        nombre = empleado["nombre"]
        guardada = guardadas.get(nombre)
        entregas = guardada["entregas"] if guardada else entregas_por_operador.get(nombre, 0)
        valor = guardada["valor_encomienda"] if guardada else empleado["valor_encomienda"]
        descuento = (
            guardada["descuento_prestamos"] if guardada
            else descuento_nomina_empleado(repository, nombre, periodo_mes)
        )
        calculo = calcular_liquidacion_semanal(
            entregas, valor, descuento, guardada["otros_descuentos"] if guardada else 0
        )
        empleados.append({
            "usuario": empleado["usuario"],
            "nombre": nombre,
            "apellidos": empleado["apellidos"],
            "cargo": empleado["cargo"],
            "entregas_reales": entregas_por_operador.get(nombre, 0),
            "observaciones": guardada["observaciones"] if guardada else "",
            "forma_pago": guardada["forma_pago"] if guardada else "",
            "liquidada": guardada is not None,
            **calculo,
        })

    return {
        "semana_inicio": semana_inicio,
        "semana_fin": domingo.isoformat(),
        "empleados": empleados,
    }


def calcular_liquidacion_laboral(
    salario_base: int,
    auxilio_transporte: int,
    fecha_ingreso: str,
    fecha_retiro: str,
    dias_prima: int | None = None,
    dias_vacaciones: int | None = None,
    indemnizacion: int = 0,
    otros_descuentos: int = 0,
) -> dict:
    """Liquidacion laboral definitiva de un empleado de nomina.

    Cesantias, intereses y prima se calculan sobre salario + auxilio de
    transporte; las vacaciones solo sobre el salario. Los dias de prima y
    de vacaciones se pueden ajustar (por defecto, todo el tiempo trabajado).
    """
    salario_base = int(salario_base or 0)
    auxilio_transporte = int(auxilio_transporte or 0)
    dias_trabajados = dias_laborales_360(fecha_ingreso, fecha_retiro)

    base_prestacional = salario_base + auxilio_transporte
    dias_prima = dias_trabajados if dias_prima is None else max(0, int(dias_prima))
    dias_vacaciones = dias_trabajados if dias_vacaciones is None else max(0, int(dias_vacaciones))

    cesantias = round(base_prestacional * dias_trabajados / DIAS_ANIO_LABORAL)
    intereses = round(cesantias * dias_trabajados * TASA_INTERESES_CESANTIAS / DIAS_ANIO_LABORAL)
    prima = round(base_prestacional * dias_prima / DIAS_ANIO_LABORAL)
    vacaciones = round(salario_base * dias_vacaciones / DIAS_VACACIONES_BASE)

    total = (
        cesantias + intereses + prima + vacaciones
        + int(indemnizacion or 0) - int(otros_descuentos or 0)
    )
    return {
        "fecha_ingreso": str(fecha_ingreso)[:10],
        "fecha_retiro": str(fecha_retiro)[:10],
        "dias_trabajados": dias_trabajados,
        "salario_base": salario_base,
        "auxilio_transporte": auxilio_transporte,
        "cesantias": cesantias,
        "intereses_cesantias": intereses,
        "prima": prima,
        "vacaciones": vacaciones,
        "indemnizacion": int(indemnizacion or 0),
        "otros_descuentos": int(otros_descuentos or 0),
        "total_pagar": total,
    }


def _escribir_hoja(sheet, titulo: str, columnas: list[str], filas: list[list], monedas: set[int]) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
    celda = sheet.cell(row=1, column=1, value=titulo)
    celda.fill = HEADER_FILL
    celda.font = TITLE_FONT
    celda.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 22

    for indice, nombre in enumerate(columnas, start=1):
        encabezado = sheet.cell(row=3, column=indice, value=nombre)
        encabezado.fill = HEADER_FILL
        encabezado.font = HEADER_FONT
        encabezado.alignment = Alignment(horizontal="center", vertical="center")

    fila_actual = 4
    for valores in filas:
        for indice, valor in enumerate(valores, start=1):
            celda = sheet.cell(row=fila_actual, column=indice, value=valor)
            if indice in monedas:
                celda.number_format = CURRENCY_FORMAT
        fila_actual += 1

    if not filas:
        sheet.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=len(columnas))
        sheet.cell(row=fila_actual, column=1, value="Sin registros").alignment = Alignment(horizontal="center")
        fila_actual += 1

    total_fila = sheet.cell(row=fila_actual, column=1, value="TOTALES")
    total_fila.fill = HEADER_FILL
    total_fila.font = HEADER_FONT
    for columna in monedas:
        letra = get_column_letter(columna)
        celda = sheet.cell(row=fila_actual, column=columna, value=f"=SUM({letra}4:{letra}{fila_actual - 1})")
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.number_format = CURRENCY_FORMAT

    for indice in range(1, len(columnas) + 1):
        sheet.column_dimensions[get_column_letter(indice)].width = 18
    sheet.column_dimensions["A"].width = 24


def generate_liquidacion_semanal_excel(
    repository: GuiaRepository, output_dir: Path, fecha: str
) -> Path:
    """Soporte de pago de la semana para los contratistas de servicios."""
    lunes, domingo = rango_semana(fecha)
    registros = repository.listar_liquidaciones_semanales(lunes.isoformat())

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"liquidacion semanal {lunes.isoformat()}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LIQUIDACION"
    _escribir_hoja(
        sheet,
        f"LIQUIDACION SEMANAL DE SERVICIOS - {lunes.isoformat()} al {domingo.isoformat()}",
        ["EMPLEADO", "ENTREGAS", "VALOR X ENCOMIENDA", "SUBTOTAL",
         "PRESTAMOS", "OTROS DESCUENTOS", "NETO A PAGAR", "FORMA DE PAGO", "OBSERVACIONES"],
        [
            [
                r["empleado"], r["entregas"], r["valor_encomienda"], r["subtotal"],
                r["descuento_prestamos"], r["otros_descuentos"], r["total_pagar"],
                r["forma_pago"], r["observaciones"],
            ]
            for r in registros
        ],
        monedas={3, 4, 5, 6, 7},
    )
    workbook.save(output_path)
    return output_path


def generate_liquidaciones_laborales_excel(
    repository: GuiaRepository, output_dir: Path, empleado: str = ""
) -> Path:
    """Historico de liquidaciones laborales definitivas."""
    registros = repository.listar_liquidaciones_laborales(empleado)

    output_dir.mkdir(parents=True, exist_ok=True)
    sufijo = f" {empleado}" if empleado else ""
    output_path = output_dir / f"liquidaciones laborales{sufijo}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LIQUIDACIONES"
    _escribir_hoja(
        sheet,
        "LIQUIDACIONES LABORALES DEFINITIVAS",
        ["EMPLEADO", "INGRESO", "RETIRO", "DIAS", "SALARIO", "AUXILIO", "CESANTIAS",
         "INT. CESANTIAS", "PRIMA", "VACACIONES", "INDEMNIZACION", "DESCUENTOS",
         "TOTAL A PAGAR", "OBSERVACIONES"],
        [
            [
                r["empleado"], r["fecha_ingreso"], r["fecha_retiro"], r["dias_trabajados"],
                r["salario_base"], r["auxilio_transporte"], r["cesantias"],
                r["intereses_cesantias"], r["prima"], r["vacaciones"],
                r["indemnizacion"], r["otros_descuentos"], r["total_pagar"],
                r["observaciones"],
            ]
            for r in registros
        ],
        monedas={5, 6, 7, 8, 9, 10, 11, 12, 13},
    )
    workbook.save(output_path)
    return output_path
