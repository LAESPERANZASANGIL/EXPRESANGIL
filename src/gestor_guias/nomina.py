"""Prestamos, adelantos de nomina y liquidacion mensual de nomina.

Reglas de negocio (definidas por la oficina):

- **Prestamo**: cobra 2% mensual sobre el saldo pendiente de capital. Las
  cuotas se pactan al momento de otorgarlo. Los abonos se aplican primero
  a los intereses causados y el resto a capital.
- **Adelanto de nomina**: no cobra interes; se descuenta de la nomina.
- **Nomina mensual**: salario proporcional a los dias trabajados, mas
  auxilio de transporte y bonificaciones, menos salud (4%), pension (4%),
  las cuotas de prestamos/adelantos y otros descuentos.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .exporter import MONTHS_ES
from .repository import GuiaRepository


TIPO_PRESTAMO = "PRESTAMO"
TIPO_ADELANTO = "ADELANTO"
TIPOS_VALIDOS = (TIPO_PRESTAMO, TIPO_ADELANTO)

# Interes mensual sobre el saldo pendiente de un prestamo.
TASA_INTERES_MENSUAL = 0.02

# Porcentajes de ley que se descuentan al empleado.
PORCENTAJE_SALUD = 0.04
PORCENTAJE_PENSION = 0.04

# Base de dias de un mes de nomina en Colombia.
DIAS_MES_NOMINA = 30

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
TOTAL_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
CURRENCY_FORMAT = '"$" #,##0'


def _periodo(fecha: str) -> tuple[int, int]:
    """De 'YYYY-MM-DD' (o 'YYYY-MM') a (anio, mes)."""
    partes = str(fecha).strip()[:7].split("-")
    return int(partes[0]), int(partes[1])


def _meses_transcurridos(desde: tuple[int, int], hasta: tuple[int, int]) -> int:
    return (hasta[0] - desde[0]) * 12 + (hasta[1] - desde[1])


def _siguiente_mes(periodo: tuple[int, int]) -> tuple[int, int]:
    anio, mes = periodo
    return (anio + 1, 1) if mes == 12 else (anio, mes + 1)


def estado_prestamo(prestamo: dict, abonos: list[dict], hasta: str = "") -> dict:
    """Calcula el estado de un prestamo o adelanto a una fecha de corte.

    Para prestamos, causa 2% mensual sobre el saldo de capital vigente al
    inicio de cada mes posterior al desembolso. Los abonos del mes se
    aplican primero a intereses y el remanente a capital.
    """
    monto = int(prestamo.get("monto", 0) or 0)
    tipo = str(prestamo.get("tipo", TIPO_PRESTAMO)).strip().upper()
    cuotas = max(1, int(prestamo.get("cuotas", 1) or 1))
    inicio = _periodo(prestamo.get("fecha", ""))
    corte = _periodo(hasta) if hasta else (date.today().year, date.today().month)

    abonos_por_mes: dict[tuple[int, int], int] = {}
    for abono in abonos:
        periodo = _periodo(abono.get("fecha", ""))
        abonos_por_mes[periodo] = abonos_por_mes.get(periodo, 0) + int(abono.get("monto", 0) or 0)

    saldo_capital = monto
    interes_pendiente = 0
    interes_causado_total = 0
    interes_del_mes_corte = 0

    periodo = inicio
    while _meses_transcurridos(periodo, corte) >= 0:
        if periodo != inicio and tipo == TIPO_PRESTAMO and saldo_capital > 0:
            interes_mes = round(saldo_capital * TASA_INTERES_MENSUAL)
            interes_pendiente += interes_mes
            interes_causado_total += interes_mes
            if periodo == corte:
                interes_del_mes_corte = interes_mes

        abono_mes = abonos_por_mes.get(periodo, 0)
        if abono_mes:
            aplicado_interes = min(abono_mes, interes_pendiente)
            interes_pendiente -= aplicado_interes
            saldo_capital = max(0, saldo_capital - (abono_mes - aplicado_interes))

        if periodo == corte:
            break
        periodo = _siguiente_mes(periodo)

    total_abonado = sum(abonos_por_mes.values())
    saldo_total = saldo_capital + interes_pendiente

    cuota_capital = round(monto / cuotas) if cuotas else monto
    if tipo == TIPO_PRESTAMO:
        cuota_sugerida = min(saldo_capital, cuota_capital) + interes_pendiente
    else:
        cuota_sugerida = min(saldo_capital, cuota_capital)

    return {
        "id": prestamo.get("id"),
        "empleado": prestamo.get("empleado", ""),
        "tipo": tipo,
        "monto": monto,
        "fecha": prestamo.get("fecha", ""),
        "forma_pago": prestamo.get("forma_pago", ""),
        "cuotas": cuotas,
        "observaciones": prestamo.get("observaciones", ""),
        "abonado": total_abonado,
        "saldo_capital": saldo_capital,
        "interes_del_mes": interes_del_mes_corte,
        "interes_pendiente": interes_pendiente,
        "interes_causado": interes_causado_total,
        "saldo_total": saldo_total,
        "cuota_sugerida": min(cuota_sugerida, saldo_total),
        "estado": "PAGADO" if saldo_total <= 0 else str(prestamo.get("estado", "ACTIVO")),
    }


def estados_prestamos(
    repository: GuiaRepository, empleado: str = "", hasta: str = "", solo_pendientes: bool = False
) -> list[dict]:
    """Estado de todos los prestamos (opcionalmente de un empleado)."""
    resultado = []
    for prestamo in repository.listar_prestamos(empleado=empleado):
        estado = estado_prestamo(prestamo, repository.listar_abonos(int(prestamo["id"])), hasta)
        if solo_pendientes and estado["saldo_total"] <= 0:
            continue
        resultado.append(estado)
    return resultado


def descuento_nomina_empleado(repository: GuiaRepository, empleado: str, periodo: str) -> int:
    """Cuanto se le debe descontar al empleado en la nomina del periodo."""
    return sum(
        estado["cuota_sugerida"]
        for estado in estados_prestamos(repository, empleado, periodo, solo_pendientes=True)
    )


def calcular_nomina_empleado(
    salario_base: int,
    dias_trabajados: int,
    auxilio_transporte: int = 0,
    bonificaciones: int = 0,
    descuento_prestamos: int = 0,
    otros_descuentos: int = 0,
) -> dict:
    """Liquida la nomina mensual de un empleado.

    Salud y pension se calculan sobre el salario devengado; el auxilio de
    transporte no es base de esos aportes.
    """
    salario_base = int(salario_base or 0)
    dias = max(0, min(int(dias_trabajados or 0), DIAS_MES_NOMINA))
    salario_devengado = round(salario_base * dias / DIAS_MES_NOMINA)
    auxilio = round(int(auxilio_transporte or 0) * dias / DIAS_MES_NOMINA)

    salud = round(salario_devengado * PORCENTAJE_SALUD)
    pension = round(salario_devengado * PORCENTAJE_PENSION)

    total_devengado = salario_devengado + auxilio + int(bonificaciones or 0)
    total_descuentos = salud + pension + int(descuento_prestamos or 0) + int(otros_descuentos or 0)

    return {
        "salario_base": salario_base,
        "dias_trabajados": dias,
        "salario_devengado": salario_devengado,
        "auxilio_transporte": auxilio,
        "bonificaciones": int(bonificaciones or 0),
        "total_devengado": total_devengado,
        "salud": salud,
        "pension": pension,
        "descuento_prestamos": int(descuento_prestamos or 0),
        "otros_descuentos": int(otros_descuentos or 0),
        "total_descuentos": total_descuentos,
        "total_pagar": total_devengado - total_descuentos,
    }


def _titulo(sheet, texto: str, ultima_columna: int, fila: int = 1) -> None:
    sheet.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima_columna)
    celda = sheet.cell(row=fila, column=1, value=texto)
    celda.fill = HEADER_FILL
    celda.font = TITLE_FONT
    celda.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[fila].height = 22


def _encabezados(sheet, columnas: list[str], fila: int) -> None:
    for indice, nombre in enumerate(columnas, start=1):
        celda = sheet.cell(row=fila, column=indice, value=nombre)
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center")


def generate_informe_prestamos_excel(
    repository: GuiaRepository, output_dir: Path, year: int, month: int
) -> Path:
    """Informe mensual de prestamos y adelantos con saldos e intereses."""
    periodo = f"{year:04d}-{month:02d}"
    estados = estados_prestamos(repository, hasta=periodo)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"prestamos y adelantos {MONTHS_ES[month]} {year}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PRESTAMOS"

    columnas = [
        "EMPLEADO", "TIPO", "FECHA", "MONTO", "FORMA DE PAGO", "CUOTAS",
        "ABONADO", "SALDO CAPITAL", "INTERES DEL MES", "INTERES PENDIENTE",
        "SALDO TOTAL", "CUOTA DEL MES", "ESTADO", "OBSERVACIONES",
    ]
    _titulo(sheet, f"PRESTAMOS Y ADELANTOS DE NOMINA - {MONTHS_ES[month].upper()} {year}", len(columnas))
    sheet.cell(row=2, column=1, value="Los prestamos causan 2% mensual sobre el saldo; los adelantos no causan interes.")
    _encabezados(sheet, columnas, 4)

    monedas = {4, 7, 8, 9, 10, 11, 12}
    fila = 5
    for estado in estados:
        valores = [
            estado["empleado"], estado["tipo"], estado["fecha"], estado["monto"],
            estado["forma_pago"], estado["cuotas"], estado["abonado"],
            estado["saldo_capital"], estado["interes_del_mes"], estado["interes_pendiente"],
            estado["saldo_total"], estado["cuota_sugerida"], estado["estado"],
            estado["observaciones"],
        ]
        for indice, valor in enumerate(valores, start=1):
            celda = sheet.cell(row=fila, column=indice, value=valor)
            if indice in monedas:
                celda.number_format = CURRENCY_FORMAT
        fila += 1

    if fila == 5:
        sheet.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(columnas))
        sheet.cell(row=fila, column=1, value="Sin prestamos ni adelantos registrados").alignment = Alignment(
            horizontal="center"
        )
        fila += 1

    total_fila = fila
    celda_total = sheet.cell(row=total_fila, column=1, value="TOTALES")
    celda_total.fill = HEADER_FILL
    celda_total.font = HEADER_FONT
    for columna in (4, 7, 8, 9, 10, 11, 12):
        letra = get_column_letter(columna)
        celda = sheet.cell(row=total_fila, column=columna, value=f"=SUM({letra}5:{letra}{total_fila - 1})")
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.number_format = CURRENCY_FORMAT

    anchos = [22, 12, 12, 14, 16, 8, 14, 15, 16, 17, 14, 14, 12, 30]
    for indice, ancho in enumerate(anchos, start=1):
        sheet.column_dimensions[get_column_letter(indice)].width = ancho

    workbook.save(output_path)
    return output_path


def generate_nomina_excel(repository: GuiaRepository, output_dir: Path, year: int, month: int) -> Path:
    """Informe mensual de nomina con devengados, deducciones y neto a pagar."""
    periodo = f"{year:04d}-{month:02d}"
    registros = repository.listar_nomina(periodo)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"nomina {MONTHS_ES[month]} {year}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NOMINA"

    columnas = [
        "EMPLEADO", "SALARIO BASE", "DIAS", "AUX. TRANSPORTE", "BONIFICACIONES",
        "SALUD (4%)", "PENSION (4%)", "PRESTAMOS", "OTROS DESCUENTOS",
        "NETO A PAGAR", "OBSERVACIONES",
    ]
    _titulo(sheet, f"NOMINA - {MONTHS_ES[month].upper()} {year}", len(columnas))
    _encabezados(sheet, columnas, 3)

    monedas = {2, 4, 5, 6, 7, 8, 9, 10}
    fila = 4
    for registro in registros:
        valores = [
            registro["empleado"], registro["salario_base"], registro["dias_trabajados"],
            registro["auxilio_transporte"], registro["bonificaciones"], registro["salud"],
            registro["pension"], registro["descuento_prestamos"], registro["otros_descuentos"],
            registro["total_pagar"], registro["observaciones"],
        ]
        for indice, valor in enumerate(valores, start=1):
            celda = sheet.cell(row=fila, column=indice, value=valor)
            if indice in monedas:
                celda.number_format = CURRENCY_FORMAT
        fila += 1

    if fila == 4:
        sheet.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(columnas))
        sheet.cell(row=fila, column=1, value="Sin nomina liquidada para este mes").alignment = Alignment(
            horizontal="center"
        )
        fila += 1

    total_fila = fila
    celda_total = sheet.cell(row=total_fila, column=1, value="TOTALES")
    celda_total.fill = HEADER_FILL
    celda_total.font = HEADER_FONT
    for columna in monedas:
        letra = get_column_letter(columna)
        celda = sheet.cell(row=total_fila, column=columna, value=f"=SUM({letra}4:{letra}{total_fila - 1})")
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.number_format = CURRENCY_FORMAT

    anchos = [22, 15, 8, 16, 16, 13, 13, 14, 17, 15, 30]
    for indice, ancho in enumerate(anchos, start=1):
        sheet.column_dimensions[get_column_letter(indice)].width = ancho

    workbook.save(output_path)
    return output_path


def generate_nomina_pdf(repository: GuiaRepository, output_dir: Path, year: int, month: int) -> Path:
    """Version PDF del informe de nomina del mes."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    periodo = f"{year:04d}-{month:02d}"
    registros = repository.listar_nomina(periodo)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"nomina {MONTHS_ES[month]} {year}.pdf"

    def pesos(valor: object) -> str:
        return f"$ {int(valor or 0):,}".replace(",", ".")

    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f"NOMINA - {MONTHS_ES[month].upper()} {year}", estilos["Title"]),
        Paragraph("Oficina Expresangil - Liquidacion mensual", estilos["Normal"]),
        Spacer(1, 0.5 * cm),
    ]

    filas = [["EMPLEADO", "SALARIO", "DIAS", "AUX.", "BONIF.", "SALUD", "PENSION", "PRESTAMOS", "NETO"]]
    for registro in registros:
        filas.append([
            str(registro["empleado"]),
            pesos(registro["salario_base"]),
            str(registro["dias_trabajados"]),
            pesos(registro["auxilio_transporte"]),
            pesos(registro["bonificaciones"]),
            pesos(registro["salud"]),
            pesos(registro["pension"]),
            pesos(registro["descuento_prestamos"]),
            pesos(registro["total_pagar"]),
        ])
    if len(filas) == 1:
        filas.append(["Sin nomina liquidada para este mes", "", "", "", "", "", "", "", ""])
    else:
        filas.append([
            "TOTAL", "", "", "", "", "", "", "",
            pesos(sum(int(r["total_pagar"] or 0) for r in registros)),
        ])

    tabla = Table(filas, colWidths=[5 * cm, 3 * cm, 1.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 3 * cm, 3 * cm])
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]
    if len(registros):
        estilo.append(("BACKGROUND", (0, len(filas) - 1), (-1, len(filas) - 1), colors.HexColor("#FFF2CC")))
        estilo.append(("FONTNAME", (0, len(filas) - 1), (-1, len(filas) - 1), "Helvetica-Bold"))
    tabla.setStyle(TableStyle(estilo))
    elementos.append(tabla)

    documento = SimpleDocTemplate(str(output_path), pagesize=landscape(letter), title=output_path.stem)
    documento.build(elementos)
    return output_path
