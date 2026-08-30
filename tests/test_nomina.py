from pathlib import Path

from gestor_guias.nomina import (
    calcular_nomina_empleado,
    descuento_nomina_empleado,
    estado_prestamo,
    estados_prestamos,
    generate_informe_prestamos_excel,
    generate_nomina_excel,
    generate_nomina_pdf,
)
from gestor_guias.repository import GuiaRepository
from openpyxl import load_workbook


def test_prestamo_causa_2_por_ciento_mensual_sobre_el_saldo() -> None:
    prestamo = {"id": 1, "empleado": "KEVIN", "tipo": "PRESTAMO", "monto": 1_000_000,
                "fecha": "2026-01-15", "cuotas": 4}

    # El mes del desembolso no causa interes.
    inicial = estado_prestamo(prestamo, [], hasta="2026-01")
    assert inicial["interes_pendiente"] == 0
    assert inicial["saldo_total"] == 1_000_000

    # Un mes despues: 2% de 1.000.000.
    un_mes = estado_prestamo(prestamo, [], hasta="2026-02")
    assert un_mes["interes_del_mes"] == 20_000
    assert un_mes["saldo_total"] == 1_020_000

    # Dos meses sin abonar: se suma otro 2% del capital vigente.
    dos_meses = estado_prestamo(prestamo, [], hasta="2026-03")
    assert dos_meses["interes_pendiente"] == 40_000
    assert dos_meses["saldo_total"] == 1_040_000


def test_abono_se_aplica_primero_a_intereses_y_luego_a_capital() -> None:
    prestamo = {"id": 1, "empleado": "KEVIN", "tipo": "PRESTAMO", "monto": 1_000_000,
                "fecha": "2026-01-15", "cuotas": 4}
    abonos = [{"fecha": "2026-02-20", "monto": 270_000}]

    estado = estado_prestamo(prestamo, abonos, hasta="2026-02")

    # 20.000 cubren el interes del mes y 250.000 bajan el capital.
    assert estado["interes_pendiente"] == 0
    assert estado["saldo_capital"] == 750_000
    assert estado["saldo_total"] == 750_000

    # El mes siguiente el interes se calcula sobre el nuevo saldo.
    siguiente = estado_prestamo(prestamo, abonos, hasta="2026-03")
    assert siguiente["interes_del_mes"] == 15_000


def test_adelanto_de_nomina_no_cobra_interes() -> None:
    adelanto = {"id": 2, "empleado": "OMAR", "tipo": "ADELANTO", "monto": 300_000,
                "fecha": "2026-01-10", "cuotas": 1}

    estado = estado_prestamo(adelanto, [], hasta="2026-04")

    assert estado["interes_pendiente"] == 0
    assert estado["saldo_total"] == 300_000
    assert estado["cuota_sugerida"] == 300_000


def test_prestamo_pagado_queda_en_cero_y_marcado() -> None:
    prestamo = {"id": 3, "empleado": "PIPE", "tipo": "PRESTAMO", "monto": 100_000,
                "fecha": "2026-01-05", "cuotas": 1}
    abonos = [{"fecha": "2026-02-05", "monto": 102_000}]

    estado = estado_prestamo(prestamo, abonos, hasta="2026-02")

    assert estado["saldo_total"] == 0
    assert estado["estado"] == "PAGADO"


def test_calcular_nomina_proporcional_con_deducciones() -> None:
    liquidacion = calcular_nomina_empleado(
        salario_base=1_500_000,
        dias_trabajados=15,
        auxilio_transporte=200_000,
        bonificaciones=100_000,
        descuento_prestamos=120_000,
        otros_descuentos=0,
    )

    assert liquidacion["salario_devengado"] == 750_000
    assert liquidacion["auxilio_transporte"] == 100_000
    assert liquidacion["total_devengado"] == 950_000
    # Salud y pension solo sobre el salario, no sobre el auxilio.
    assert liquidacion["salud"] == 30_000
    assert liquidacion["pension"] == 30_000
    assert liquidacion["total_descuentos"] == 180_000
    assert liquidacion["total_pagar"] == 770_000


def test_nomina_mes_completo_no_prorratea() -> None:
    liquidacion = calcular_nomina_empleado(salario_base=1_300_000, dias_trabajados=30)

    assert liquidacion["salario_devengado"] == 1_300_000
    assert liquidacion["total_pagar"] == 1_300_000 - 52_000 * 2


def test_descuento_de_nomina_suma_las_cuotas_del_mes(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")
    repository.crear_prestamo("KEVIN", "PRESTAMO", 1_000_000, "2026-01-15", "EFECTIVO", cuotas=4)
    repository.crear_prestamo("KEVIN", "ADELANTO", 200_000, "2026-02-01", "NEQUI", cuotas=1)
    repository.crear_prestamo("OMAR", "ADELANTO", 500_000, "2026-02-01", "NEQUI", cuotas=1)

    # Cuota del prestamo: 250.000 de capital + 20.000 de interes; mas el adelanto.
    assert descuento_nomina_empleado(repository, "KEVIN", "2026-02") == 470_000
    assert descuento_nomina_empleado(repository, "OMAR", "2026-02") == 500_000


def test_prestamo_anulado_no_aparece_ni_descuenta(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")
    prestamo_id = repository.crear_prestamo("KEVIN", "PRESTAMO", 400_000, "2026-01-10", cuotas=2)

    assert repository.anular_prestamo(prestamo_id) is True
    assert estados_prestamos(repository, hasta="2026-02") == []
    assert descuento_nomina_empleado(repository, "KEVIN", "2026-02") == 0


def test_informe_de_prestamos_y_nomina_se_generan(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")
    repository.crear_prestamo("KEVIN", "PRESTAMO", 1_000_000, "2026-01-15", "EFECTIVO", cuotas=4)
    repository.guardar_nomina(
        "2026-02",
        "KEVIN",
        calcular_nomina_empleado(1_500_000, 30, 200_000, 0, 270_000, 0),
    )

    ruta_prestamos = generate_informe_prestamos_excel(repository, tmp_path, 2026, 2)
    hoja = load_workbook(ruta_prestamos)["PRESTAMOS"]
    empleados = [hoja.cell(row=f, column=1).value for f in range(5, hoja.max_row + 1)]
    assert "KEVIN" in empleados

    ruta_nomina = generate_nomina_excel(repository, tmp_path, 2026, 2)
    hoja_nomina = load_workbook(ruta_nomina)["NOMINA"]
    assert hoja_nomina.cell(row=4, column=1).value == "KEVIN"

    ruta_pdf = generate_nomina_pdf(repository, tmp_path, 2026, 2)
    assert ruta_pdf.exists() and ruta_pdf.stat().st_size > 0


def test_guardar_y_listar_nomina_actualiza_el_registro(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")

    repository.guardar_nomina("2026-02", "OMAR", calcular_nomina_empleado(1_000_000, 30))
    repository.guardar_nomina("2026-02", "OMAR", calcular_nomina_empleado(1_200_000, 30))

    registros = repository.listar_nomina("2026-02")
    assert len(registros) == 1
    assert registros[0]["salario_base"] == 1_200_000
