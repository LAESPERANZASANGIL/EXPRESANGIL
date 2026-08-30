from datetime import date
from pathlib import Path

import pandas as pd

from gestor_guias.liquidaciones import (
    calcular_liquidacion_laboral,
    calcular_liquidacion_semanal,
    dias_laborales_360,
    generate_liquidacion_semanal_excel,
    generate_liquidaciones_laborales_excel,
    inicio_de_semana,
    preparar_semana,
    rango_semana,
)
from gestor_guias.repository import GuiaRepository
from openpyxl import load_workbook


def build_dataframe(guia: str, operador: str = "") -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "PLANILLA": "1", "SERVICIO": "", "GUIA": guia, "UNID": "1",
                "TIPO DE SERVICIO": "PT", "DESTINATARIO": "Persona",
                "MUNICIPIO": "SAN GIL", "VALOR": "10000", "OPERADOR": operador,
                "ESTADO": "", "CAUSAL": "", "F_INGRESO": "2026-03-02 00:00:00",
                "F_ENTREGA": "",
            }
        ]
    )


def test_la_semana_va_de_lunes_a_domingo() -> None:
    # 2026-03-04 es miercoles.
    assert inicio_de_semana("2026-03-04") == date(2026, 3, 2)
    assert rango_semana("2026-03-04") == (date(2026, 3, 2), date(2026, 3, 8))
    # Un domingo sigue perteneciendo a la semana que empezo el lunes.
    assert inicio_de_semana("2026-03-08") == date(2026, 3, 2)


def test_liquidacion_semanal_paga_por_encomienda_y_descuenta() -> None:
    calculo = calcular_liquidacion_semanal(
        entregas=120, valor_encomienda=1_500, descuento_prestamos=50_000, otros_descuentos=10_000
    )

    assert calculo["subtotal"] == 180_000
    assert calculo["total_pagar"] == 120_000


def test_preparar_semana_cuenta_entregas_reales_del_contratista(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")
    repository.crear_operador("pipe", "hash", "PIPE")
    repository.actualizar_datos_empleado(
        "pipe", {"tipo_contrato": "SERVICIOS", "valor_encomienda": 1_200, "fecha_ingreso": "2026-01-01"}
    )
    repository.crear_operador("omar", "hash", "OMAR")
    repository.actualizar_datos_empleado("omar", {"tipo_contrato": "NOMINA", "salario_base": 1_400_000})

    # Tres entregas de PIPE dentro de la semana y una fuera.
    for indice, entrega in enumerate(["2026-03-02", "2026-03-05", "2026-03-08", "2026-03-09"]):
        guia = f"10{indice}"
        repository.save_consolidated(build_dataframe(guia))
        repository.update_guide_details(
            guia=guia, planilla="1", destinatario="Persona", direccion="", municipio="SAN GIL",
            valor="10000", operador="PIPE", estado="E", causal="", entrega=entrega,
        )

    semana = preparar_semana(repository, "2026-03-04")

    assert semana["semana_inicio"] == "2026-03-02"
    assert semana["semana_fin"] == "2026-03-08"
    # Solo los de SERVICIOS entran en la liquidacion semanal.
    assert [e["nombre"] for e in semana["empleados"]] == ["PIPE"]
    pipe = semana["empleados"][0]
    assert pipe["entregas"] == 3
    assert pipe["valor_encomienda"] == 1_200
    assert pipe["subtotal"] == 3_600


def test_preparar_semana_descuenta_prestamos_pendientes(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")
    repository.crear_operador("pipe", "hash", "PIPE")
    repository.actualizar_datos_empleado("pipe", {"tipo_contrato": "SERVICIOS", "valor_encomienda": 1_000})
    repository.crear_prestamo("PIPE", "ADELANTO", 80_000, "2026-03-01", "EFECTIVO", cuotas=1)

    semana = preparar_semana(repository, "2026-03-04")

    assert semana["empleados"][0]["descuento_prestamos"] == 80_000


def test_dias_laborales_con_convencion_de_360() -> None:
    assert dias_laborales_360("2026-01-01", "2027-01-01") == 360
    assert dias_laborales_360("2026-01-01", "2026-07-01") == 180
    assert dias_laborales_360("2026-03-01", "2026-03-16") == 15
    # Una fecha de retiro anterior al ingreso no produce dias negativos.
    assert dias_laborales_360("2026-05-01", "2026-01-01") == 0


def test_liquidacion_laboral_de_un_anio_completo() -> None:
    liquidacion = calcular_liquidacion_laboral(
        salario_base=1_300_000,
        auxilio_transporte=200_000,
        fecha_ingreso="2026-01-01",
        fecha_retiro="2027-01-01",
    )

    assert liquidacion["dias_trabajados"] == 360
    # Un anio completo: cesantias equivalen a un mes de salario mas auxilio.
    assert liquidacion["cesantias"] == 1_500_000
    assert liquidacion["intereses_cesantias"] == 180_000
    assert liquidacion["prima"] == 1_500_000
    # Vacaciones: 15 dias de salario por anio.
    assert liquidacion["vacaciones"] == 650_000
    assert liquidacion["total_pagar"] == 3_830_000


def test_liquidacion_laboral_ajusta_dias_de_prima_y_descuenta() -> None:
    liquidacion = calcular_liquidacion_laboral(
        salario_base=1_200_000,
        auxilio_transporte=0,
        fecha_ingreso="2026-01-01",
        fecha_retiro="2026-07-01",
        dias_prima=180,
        indemnizacion=500_000,
        otros_descuentos=100_000,
    )

    assert liquidacion["dias_trabajados"] == 180
    assert liquidacion["cesantias"] == 600_000
    assert liquidacion["prima"] == 600_000
    assert liquidacion["indemnizacion"] == 500_000
    assert liquidacion["total_pagar"] == (
        600_000 + liquidacion["intereses_cesantias"] + 600_000
        + liquidacion["vacaciones"] + 500_000 - 100_000
    )


def test_liquidaciones_quedan_almacenadas_como_soporte(tmp_path: Path) -> None:
    repository = GuiaRepository(tmp_path / "guias.db")

    repository.guardar_liquidacion_semanal(
        "2026-03-02",
        "PIPE",
        calcular_liquidacion_semanal(100, 1_200, 50_000, 0) | {"forma_pago": "NEQUI"},
    )
    # Volver a guardar la misma semana actualiza, no duplica.
    repository.guardar_liquidacion_semanal(
        "2026-03-02", "PIPE", calcular_liquidacion_semanal(110, 1_200, 50_000, 0)
    )
    semanales = repository.listar_liquidaciones_semanales("2026-03-02")
    assert len(semanales) == 1
    assert semanales[0]["entregas"] == 110
    assert semanales[0]["registrada_en"] != ""

    laboral = calcular_liquidacion_laboral(1_300_000, 200_000, "2026-01-01", "2027-01-01")
    repository.guardar_liquidacion_laboral({**laboral, "empleado": "KEVIN", "observaciones": "retiro"})
    laborales = repository.listar_liquidaciones_laborales("KEVIN")
    assert len(laborales) == 1
    assert laborales[0]["total_pagar"] == laboral["total_pagar"]

    ruta_semanal = generate_liquidacion_semanal_excel(repository, tmp_path, "2026-03-04")
    hoja = load_workbook(ruta_semanal)["LIQUIDACION"]
    assert hoja.cell(row=4, column=1).value == "PIPE"

    ruta_laboral = generate_liquidaciones_laborales_excel(repository, tmp_path)
    hoja_laboral = load_workbook(ruta_laboral)["LIQUIDACIONES"]
    assert hoja_laboral.cell(row=4, column=1).value == "KEVIN"
